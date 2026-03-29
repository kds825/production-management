"""재공실사 Excel 템플릿 생성 — 드롭다운 validation 포함"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


# 컬럼 정의
COLUMNS = [
    "공정",
    "전압구분",
    "재질",
    "품명",
    "규격",
    "길이(M)",
    "개수",
    "총량(M)",
    "선심색상",
    "소선경",
    "가닥수",
    "상태",
]

# 드롭다운 값 — KBI 현장 기준
DROPDOWNS = {
    "공정": ["연선재고", "절연재고", "연합재고", "완제품"],
    "전압구분": ["저압", "고압"],
    "재질": ["CU", "AL"],
    "품명": ["N/A", "TFR-CV(WB)", "TFR-8", "TFR-8 고내화", "URD"],
    "규격": [
        "16SQ",
        "25SQ",
        "35SQ",
        "50SQ",
        "70SQ",
        "95SQ",
        "120SQ",
        "150SQ",
        "185SQ",
        "200SQ",
        "240SQ",
        "300SQ",
        "400SQ",
        "500kcmil",
        "750kcmil",
        "1000kcmil",
        "1250kcmil",
    ],
    "선심색상": ["흑", "갈", "회", "청", "녹/황", "흑/적", "흑/갈/회/청"],
    "상태": ["예상", "실적"],
}

# 스타일
_HEADER_FONT = Font(bold=True, size=10, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="006400")  # 진녹색 (이미지 기준)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")  # 연노랑 (입력 셀)
_CENTER = Alignment(horizontal="center", vertical="center")

COL_WIDTHS = [14, 12, 8, 16, 14, 12, 8, 12, 12, 10, 10, 8]


def generate_wip_template() -> BytesIO:
    """재공실사 데이터 입력용 Excel 템플릿 생성. 드롭다운 validation 포함."""
    wb = Workbook()
    ws = wb.active
    ws.title = "재공실사"

    # ── 헤더 ──────────────────────────────────────────────────────────────
    for col_idx, col_name in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _CENTER

    # ── 열 너비 ──────────────────────────────────────────────────────────
    for i, width in enumerate(COL_WIDTHS, 1):
        from openpyxl.utils import get_column_letter

        ws.column_dimensions[get_column_letter(i)].width = width

    # ── 데이터 행 서식 (50행 준비) ────────────────────────────────────────
    for r in range(2, 52):
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _THIN_BORDER
            # 입력 컬럼 배경색
            col_name = COLUMNS[c - 1]
            if col_name in ("길이(M)", "개수", "소선경", "가닥수"):
                cell.fill = _INPUT_FILL

    # ── 드롭다운 Validation ──────────────────────────────────────────────
    max_row = 51
    for col_name, values in DROPDOWNS.items():
        if col_name not in COLUMNS:
            continue
        col_idx = COLUMNS.index(col_name) + 1
        col_letter = get_column_letter(col_idx)
        formula = ",".join(values)

        dv = DataValidation(
            type="list",
            formula1=f'"{formula}"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="입력 오류",
            error=f"허용된 값: {formula}",
        )
        dv.sqref = f"{col_letter}2:{col_letter}{max_row}"
        ws.add_data_validation(dv)

    # ── 총량(M) 자동 계산 수식 ────────────────────────────────────────────
    total_col = COLUMNS.index("총량(M)") + 1
    len_col = get_column_letter(COLUMNS.index("길이(M)") + 1)
    cnt_col = get_column_letter(COLUMNS.index("개수") + 1)
    for r in range(2, max_row + 1):
        ws.cell(row=r, column=total_col).value = (
            f'=IF(AND({len_col}{r}<>"",{cnt_col}{r}<>""),{len_col}{r}*{cnt_col}{r},"")'
        )

    # 헤더 고정
    ws.freeze_panes = "A2"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
