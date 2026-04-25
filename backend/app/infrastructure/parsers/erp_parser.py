"""ERP 수주 파일(.xls) 파싱 — 진행/대기 시트에서 SalesOrder 레코드 생성"""

import re
from datetime import datetime

import xlrd
from sqlalchemy.orm import Session

from app.infrastructure.models.sales_order import SalesOrder
from app.infrastructure.models.item_master import ItemMaster


def parse_erp_file(file_content: bytes, run_label: str, db: Session) -> dict:
    """
    ERP .xls 파일을 파싱하여 sales_order 테이블에 저장.
    기존 데이터를 전체 삭제한 뒤 파일의 모든 행을 신규 INSERT한다.

    Args:
        file_content: .xls 파일의 raw bytes
        run_label:    계획 실행 식별자 (e.g. "20250325_103045")
        db:           SQLAlchemy 세션 (commit은 호출부 책임)

    Returns:
        {
            "total":    int,       # 처리된 전체 행 수
            "진행":     int,       # 진행 시트 행 수
            "대기":     int,       # 대기 시트 행 수
            "외주_제외": int,       # is_outsourced=True 로 마킹된 행 수
            "inserted": int,       # 삽입 건수
            "warnings": list[str], # 비치명적 이슈 메시지
        }
    """
    workbook = xlrd.open_workbook(file_contents=file_content)
    result: dict = {
        "total": 0,
        "진행": 0,
        "대기": 0,
        "외주_제외": 0,
        "inserted": 0,
        "warnings": [],
    }

    # ── 기존 sales_order 전체 삭제 ───────────────────────────────────────────
    deleted = db.query(SalesOrder).delete(synchronize_session=False)
    if deleted:
        result["warnings"].append(f"기존 수주 {deleted}건 삭제 후 재적재")

    # ── item_master 룩업: (product_group, voltage) → item_code ───────────────
    items = db.query(ItemMaster).all()
    _item_by_pg_volt: dict[tuple[str, str], str] = {}
    _item_by_pg: dict[str, str] = {}
    for it in items:
        pg = (it.product_group or "").strip()
        vt = (it.voltage or "").strip()
        if pg and vt and (pg, vt) not in _item_by_pg_volt:
            _item_by_pg_volt[(pg, vt)] = it.item_code
        if pg and pg not in _item_by_pg:
            _item_by_pg[pg] = it.item_code

    def _resolve_item_code(product_group: str, voltage: str | None) -> str | None:
        pg = (product_group or "").strip()
        vt = (voltage or "").strip()
        if not pg:
            return None
        return _item_by_pg_volt.get((pg, vt)) or _item_by_pg.get(pg)

    # order_line: 전체 INSERT 순번 (1부터 시작)
    line_counter = 0

    for sheet_name in ["진행", "대기"]:
        if sheet_name not in workbook.sheet_names():
            result["warnings"].append(f"시트 '{sheet_name}' 없음 — 건너뜀")
            continue

        sheet = workbook.sheet_by_name(sheet_name)

        # ── 헤더 행 탐지: "수주번호"를 포함한 셀이 있는 첫 번째 행 ──────────────
        header_row_idx = _find_header_row(sheet, keyword="수주번호", max_scan=25)
        if header_row_idx is None:
            result["warnings"].append(
                f"시트 '{sheet_name}': 헤더 행 탐지 실패 — 건너뜀"
            )
            continue

        # ── 컬럼명 → 인덱스 맵 구성 ────────────────────────────────────────────
        headers = _build_header_map(sheet, header_row_idx)
        order_id_col_idx = headers.get("수주번호")

        # ── 데이터 행 파싱 ───────────────────────────────────────────────────────
        for r in range(header_row_idx + 1, sheet.nrows):
            check_col = order_id_col_idx if order_id_col_idx is not None else 0
            if not str(sheet.cell_value(r, check_col)).strip():
                continue

            try:
                order_id = _get_str(sheet, r, headers, "수주번호")
                if not order_id:
                    continue

                outsource_plan = _get_str(sheet, r, headers, "외주계획") or ""
                is_outsourced = outsource_plan.strip().upper() == "Y"

                due_date_raw = _get_cell(sheet, r, headers, "납품일")
                due_date = _parse_date(due_date_raw, workbook.datemode)

                spec_raw = _get_str(sheet, r, headers, "규격") or ""
                core_count = _extract_core_count(spec_raw)

                drum_length_m = _get_num(sheet, r, headers, "(수주)조장(M)")
                product_group = _get_str(sheet, r, headers, "제품군") or ""
                voltage = _get_str(sheet, r, headers, "전압")
                item_code = _resolve_item_code(product_group, voltage)

                line_counter += 1
                order = SalesOrder(
                    order_id=order_id,
                    order_line=line_counter,
                    order_status=sheet_name,
                    item_code=item_code,
                    product_group=product_group,
                    voltage=voltage,
                    spec_raw=spec_raw,
                    customer_name=_get_str(sheet, r, headers, "거래처명"),
                    due_date=due_date,
                    due_type="출하기준",
                    drum_length_m=drum_length_m,
                    drum_count=_get_int(sheet, r, headers, "(수주)개수(ea)"),
                    ordered_qty_m=_get_num(sheet, r, headers, "(수주)수량(M)"),
                    self_plan_qty_m=_get_num(sheet, r, headers, "(자체)조장"),
                    unit_price_krw=_get_num(sheet, r, headers, "원화단가"),
                    amount_krw=_get_num(sheet, r, headers, "원화금액"),
                    cu_weight_kg=_get_num(sheet, r, headers, "CU량"),
                    al_weight_kg=_get_num(sheet, r, headers, "AL량"),
                    core_count=core_count,
                    core_colors=_get_str(sheet, r, headers, "선심색상"),
                    sheath_color=_get_str(sheet, r, headers, "색상"),
                    neutral_wire=_get_str(sheet, r, headers, "중성선"),
                    is_outsourced=is_outsourced,
                    run_label=run_label,
                )
                db.add(order)
                result["inserted"] += 1

                if is_outsourced:
                    result["외주_제외"] += 1

                result[sheet_name] += 1
                result["total"] += 1

            except Exception as e:
                result["warnings"].append(f"[{sheet_name}] 행 {r + 1}: {e}")

    db.flush()
    return result


def parse_erp_file_incremental(
    file_content: bytes, run_label: str, db: Session
) -> dict:
    """ERP 파일을 증분(incremental) 파싱하여 기존 SalesOrder를 유지하고 새 수주만 추가.

    기존 parse_erp_file()과의 차이:
    - 기존 sales_order를 삭제하지 않음
    - .xlsx(openpyxl) 우선, .xls(xlrd) 폴백
    - 시트명 "진행"/"대기" 고정이 아닌, "수주번호" 헤더가 있는 모든 시트 자동 탐지
    - 중복 감지는 (order_id, drum_length_m) 쌍 기준 — 같은 수주번호의 다른 드럼은 별도 행으로 삽입
    - line_counter는 기존 max(order_line) + 1부터 시작

    Args:
        file_content: .xlsx 또는 .xls 파일의 raw bytes
        run_label:    계획 실행 식별자 (기존 run_label 재사용)
        db:           SQLAlchemy 세션 (commit은 호출부 책임)

    Returns:
        {
            "total": int, "inserted": int, "skipped_dup": int,
            "외주_제외": int, "warnings": list[str],
        }
    """
    from io import BytesIO

    result: dict = {
        "total": 0,
        "inserted": 0,
        "skipped_dup": 0,
        "외주_제외": 0,
        "warnings": [],
    }

    # ── 기존 DB의 (order_id, drum_length_m) 쌍 로드 (DB 중복 감지용) ────────
    # 같은 order_id라도 규격/길이가 다르면 별개 수주이므로 (order_id, drum_length_m) 쌍으로 비교.
    # DB에 이미 존재하는 쌍만 스킵하고, 파일 내 동일 쌍은 모두 삽입한다
    # (같은 order_id + 같은 길이의 서로 다른 드럼은 별도 order_line으로 삽입됨).
    existing_pairs: set[tuple[str, float | None]] = {
        (oid, float(length) if length is not None else None)
        for (oid, length) in db.query(SalesOrder.order_id, SalesOrder.drum_length_m)
        .filter(SalesOrder.run_label == run_label)
        .all()
    }

    # ── line_counter: 기존 max(order_line) + 1부터 시작 ───────────────────────
    from sqlalchemy import func

    max_line = (
        db.query(func.max(SalesOrder.order_line))
        .filter(SalesOrder.run_label == run_label)
        .scalar()
    )
    line_counter = max_line or 0

    # ── item_master 룩업: (product_group, voltage) → item_code ───────────────
    items = db.query(ItemMaster).all()
    _item_by_pg_volt: dict[tuple[str, str], str] = {}
    _item_by_pg: dict[str, str] = {}
    for it in items:
        pg = (it.product_group or "").strip()
        vt = (it.voltage or "").strip()
        if pg and vt and (pg, vt) not in _item_by_pg_volt:
            _item_by_pg_volt[(pg, vt)] = it.item_code
        if pg and pg not in _item_by_pg:
            _item_by_pg[pg] = it.item_code

    def _resolve_item_code(product_group: str, voltage: str | None) -> str | None:
        pg = (product_group or "").strip()
        vt = (voltage or "").strip()
        if not pg:
            return None
        return _item_by_pg_volt.get((pg, vt)) or _item_by_pg.get(pg)

    # ── .xlsx 시도 (openpyxl) → 실패 시 .xls 폴백 (xlrd) ────────────────────
    parsed = False
    try:
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(file_content), data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            header_row_idx = _find_header_row_openpyxl(
                ws, keyword="수주번호", max_scan=25
            )
            if header_row_idx is None:
                continue  # 이 시트에는 수주 데이터 없음

            headers = _build_header_map_openpyxl(ws, header_row_idx)
            order_id_col = headers.get("수주번호")

            for r in range(header_row_idx + 1, ws.max_row + 1):
                check_col = order_id_col if order_id_col is not None else 1
                cell_val = ws.cell(r, check_col).value
                if not str(cell_val or "").strip():
                    continue

                try:
                    order_id = _get_str_openpyxl(ws, r, headers, "수주번호")
                    if not order_id:
                        continue

                    outsource_plan = _get_str_openpyxl(ws, r, headers, "외주계획") or ""
                    is_outsourced = outsource_plan.strip().upper() == "Y"

                    due_date_raw = _get_cell_openpyxl(ws, r, headers, "납품일")
                    due_date = _parse_date_openpyxl(due_date_raw)

                    spec_raw = _get_str_openpyxl(ws, r, headers, "규격") or ""
                    core_count = _extract_core_count(spec_raw)

                    drum_length_m = _get_num_openpyxl(ws, r, headers, "(수주)조장(M)")
                    product_group = _get_str_openpyxl(ws, r, headers, "제품군") or ""
                    voltage = _get_str_openpyxl(ws, r, headers, "전압")
                    item_code = _resolve_item_code(product_group, voltage)

                    # 중복 감지: DB에 이미 존재하는 (order_id, drum_length_m) 쌍이면 건너뜀.
                    # 같은 order_id라도 길이/규격이 다르면 별개 수주이므로 삽입.
                    # 파일 내 동일 order_id+길이 행은 별개 드럼이므로 모두 삽입.
                    pair = (order_id, drum_length_m)
                    if pair in existing_pairs:
                        result["skipped_dup"] += 1
                        continue

                    line_counter += 1
                    order = SalesOrder(
                        order_id=order_id,
                        order_line=line_counter,
                        order_status=sheet_name,
                        item_code=item_code,
                        product_group=product_group,
                        voltage=voltage,
                        spec_raw=spec_raw,
                        customer_name=_get_str_openpyxl(ws, r, headers, "거래처명"),
                        due_date=due_date,
                        due_type="출하기준",
                        drum_length_m=drum_length_m,
                        drum_count=_get_int_openpyxl(ws, r, headers, "(수주)개수(ea)"),
                        ordered_qty_m=_get_num_openpyxl(
                            ws, r, headers, "(수주)수량(M)"
                        ),
                        self_plan_qty_m=_get_num_openpyxl(ws, r, headers, "(자체)조장"),
                        unit_price_krw=_get_num_openpyxl(ws, r, headers, "원화단가"),
                        amount_krw=_get_num_openpyxl(ws, r, headers, "원화금액"),
                        cu_weight_kg=_get_num_openpyxl(ws, r, headers, "CU량"),
                        al_weight_kg=_get_num_openpyxl(ws, r, headers, "AL량"),
                        core_count=core_count,
                        core_colors=_get_str_openpyxl(ws, r, headers, "선심색상"),
                        sheath_color=_get_str_openpyxl(ws, r, headers, "색상"),
                        neutral_wire=_get_str_openpyxl(ws, r, headers, "중성선"),
                        is_outsourced=is_outsourced,
                        run_label=run_label,
                    )
                    db.add(order)
                    result["inserted"] += 1

                    if is_outsourced:
                        result["외주_제외"] += 1

                    result["total"] += 1

                except Exception as e:
                    result["warnings"].append(f"[{sheet_name}] 행 {r}: {e}")

        parsed = True
    except Exception:
        pass  # .xlsx 파싱 실패 → .xls 폴백

    if not parsed:
        # .xls 폴백 (xlrd) — 기존 parse_erp_file과 동일한 xlrd 로직 사용
        try:
            workbook = xlrd.open_workbook(file_contents=file_content)
        except Exception as exc:
            raise ValueError(f"xlsx/xls 모두 파싱 실패: {exc}") from exc

        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)
            header_row_idx = _find_header_row(sheet, keyword="수주번호", max_scan=25)
            if header_row_idx is None:
                continue

            headers = _build_header_map(sheet, header_row_idx)
            order_id_col_idx = headers.get("수주번호")

            for r in range(header_row_idx + 1, sheet.nrows):
                check_col = order_id_col_idx if order_id_col_idx is not None else 0
                if not str(sheet.cell_value(r, check_col)).strip():
                    continue

                try:
                    order_id = _get_str(sheet, r, headers, "수주번호")
                    if not order_id:
                        continue

                    outsource_plan = _get_str(sheet, r, headers, "외주계획") or ""
                    is_outsourced = outsource_plan.strip().upper() == "Y"

                    due_date_raw = _get_cell(sheet, r, headers, "납품일")
                    due_date = _parse_date(due_date_raw, workbook.datemode)

                    spec_raw = _get_str(sheet, r, headers, "규격") or ""
                    core_count = _extract_core_count(spec_raw)

                    drum_length_m = _get_num(sheet, r, headers, "(수주)조장(M)")
                    product_group = _get_str(sheet, r, headers, "제품군") or ""
                    voltage = _get_str(sheet, r, headers, "전압")
                    item_code = _resolve_item_code(product_group, voltage)

                    # 중복 감지: DB에 이미 존재하는 (order_id, drum_length_m) 쌍이면 건너뜀
                    pair = (order_id, drum_length_m)
                    if pair in existing_pairs:
                        result["skipped_dup"] += 1
                        continue

                    line_counter += 1
                    order = SalesOrder(
                        order_id=order_id,
                        order_line=line_counter,
                        order_status=sheet_name,
                        item_code=item_code,
                        product_group=product_group,
                        voltage=voltage,
                        spec_raw=spec_raw,
                        customer_name=_get_str(sheet, r, headers, "거래처명"),
                        due_date=due_date,
                        due_type="출하기준",
                        drum_length_m=drum_length_m,
                        drum_count=_get_int(sheet, r, headers, "(수주)개수(ea)"),
                        ordered_qty_m=_get_num(sheet, r, headers, "(수주)수량(M)"),
                        self_plan_qty_m=_get_num(sheet, r, headers, "(자체)조장"),
                        unit_price_krw=_get_num(sheet, r, headers, "원화단가"),
                        amount_krw=_get_num(sheet, r, headers, "원화금액"),
                        cu_weight_kg=_get_num(sheet, r, headers, "CU량"),
                        al_weight_kg=_get_num(sheet, r, headers, "AL량"),
                        core_count=core_count,
                        core_colors=_get_str(sheet, r, headers, "선심색상"),
                        sheath_color=_get_str(sheet, r, headers, "색상"),
                        neutral_wire=_get_str(sheet, r, headers, "중성선"),
                        is_outsourced=is_outsourced,
                        run_label=run_label,
                    )
                    db.add(order)
                    result["inserted"] += 1

                    if is_outsourced:
                        result["외주_제외"] += 1

                    result["total"] += 1

                except Exception as e:
                    result["warnings"].append(f"[{sheet_name}] 행 {r + 1}: {e}")

    if result["inserted"] == 0 and result["skipped_dup"] == 0:
        result["warnings"].append(
            "수주 데이터를 찾을 수 없습니다 — 수주번호 헤더가 있는 시트 없음"
        )

    db.flush()
    return result


# ── openpyxl 전용 헬퍼 (증분 파서용) ──────────────────────────────────────────


def _find_header_row_openpyxl(ws, keyword: str, max_scan: int) -> int | None:
    """openpyxl 워크시트에서 헤더 행을 탐지한다.

    keyword를 포함하면서 다른 컬럼 키워드도 함께 있는 행을 찾는다.
    반환값은 1-based 행 번호 (openpyxl 규약).
    """
    confirm_keywords = {"전압", "제품군", "규격", "거래처", "납품일"}
    for r in range(1, min(max_scan + 1, ws.max_row + 1)):
        row_texts = [
            str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)
        ]
        has_keyword = any(
            keyword == t or keyword == t.rstrip(":").strip() for t in row_texts
        )
        if not has_keyword:
            continue
        confirm_count = sum(
            1 for ck in confirm_keywords if any(ck in t for t in row_texts)
        )
        if confirm_count >= 2:
            return r
    return None


def _build_header_map_openpyxl(ws, header_row: int) -> dict[str, int]:
    """openpyxl 워크시트에서 {컬럼명: 열번호(1-based)} 딕셔너리를 반환."""
    headers: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        val = str(ws.cell(header_row, c).value or "").strip()
        val = val.replace("\n", " ").replace("\r", " ").strip()
        if val and val not in headers:
            headers[val] = c
    return headers


def _get_cell_openpyxl(ws, row: int, headers: dict[str, int], col_name: str):
    """openpyxl 워크시트에서 셀 값을 반환. 부분 일치 폴백 포함."""
    col_idx = headers.get(col_name)
    if col_idx is None:
        norm = col_name.replace(" ", "").replace("\n", "").lower()
        for h, idx in headers.items():
            h_norm = h.replace(" ", "").replace("\n", "").lower()
            if norm in h_norm or h_norm in norm:
                col_idx = idx
                break
    if col_idx is None:
        return None
    return ws.cell(row, col_idx).value


def _get_str_openpyxl(
    ws, row: int, headers: dict[str, int], col_name: str
) -> str | None:
    """openpyxl 셀 값을 문자열로 반환."""
    val = _get_cell_openpyxl(ws, row, headers, col_name)
    if val is None:
        return None
    if isinstance(val, float) and val == int(val):
        s = str(int(val))
    else:
        s = str(val).strip()
    return s if s else None


def _get_num_openpyxl(
    ws, row: int, headers: dict[str, int], col_name: str
) -> float | None:
    """openpyxl 셀 값을 float으로 반환."""
    val = _get_cell_openpyxl(ws, row, headers, col_name)
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _get_int_openpyxl(
    ws, row: int, headers: dict[str, int], col_name: str
) -> int | None:
    """openpyxl 셀 값을 int로 반환."""
    val = _get_num_openpyxl(ws, row, headers, col_name)
    return int(val) if val is not None else None


def _parse_date_openpyxl(val):
    """openpyxl 날짜 값을 date로 변환.

    openpyxl은 날짜 셀을 datetime 객체로 자동 변환하므로 xlrd와 다르게 처리한다.
    """
    if val is None or val == "":
        return None

    # openpyxl이 이미 datetime으로 변환한 경우
    if isinstance(val, datetime):
        return val.date()

    # date 객체인 경우
    from datetime import date as _date

    if isinstance(val, _date):
        return val

    # 숫자형 (YYYYMMDD)
    if isinstance(val, (int, float)):
        int_val = int(val)
        if 19000101 <= int_val <= 29991231:
            try:
                return datetime.strptime(str(int_val), "%Y%m%d").date()
            except ValueError:
                pass
        return None

    # 문자열형
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


# ── 내부 헬퍼 ─────────────────────────────────────────────────────────────────


def _find_header_row(sheet, keyword: str, max_scan: int) -> int | None:
    """헤더 행 탐지 — keyword를 포함하면서 다른 컬럼 키워드도 함께 있는 행.
    메타데이터 행("수주번호 :")과 실제 헤더 행을 구분."""
    confirm_keywords = {"전압", "제품군", "규격", "거래처", "납품일"}
    for r in range(min(max_scan, sheet.nrows)):
        row_texts = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
        has_keyword = any(
            keyword == t or keyword == t.rstrip(":").strip() for t in row_texts
        )
        if not has_keyword:
            continue
        # 다른 컬럼 키워드 최소 2개 이상 존재하면 헤더 행으로 인정
        confirm_count = sum(
            1 for ck in confirm_keywords if any(ck in t for t in row_texts)
        )
        if confirm_count >= 2:
            return r
    return None


def _build_header_map(sheet, header_row_idx: int) -> dict[str, int]:
    """헤더 행을 순회하여 {컬럼명: 열인덱스} 딕셔너리를 반환.

    동일 컬럼명이 여러 번 등장하면 첫 번째 위치를 사용한다.
    """
    headers: dict[str, int] = {}
    for c in range(sheet.ncols):
        val = str(sheet.cell_value(header_row_idx, c)).strip()
        # 줄바꿈을 공백으로 치환 (ERP 엑셀의 셀 내 줄바꿈 처리)
        val = val.replace("\n", " ").replace("\r", " ").strip()
        if val and val not in headers:
            headers[val] = c
    return headers


def _get_cell(sheet, row: int, headers: dict[str, int], col_name: str):
    """헤더 맵에서 col_name을 찾아 셀 값을 반환.

    정확히 일치하는 키가 없으면 col_name을 부분 문자열로 포함하는 첫 번째 키로
    폴백한다 (e.g. "(수주)조장(M)" vs 공백 포함 변형).
    컬럼 자체가 없으면 None 반환.
    """
    col_idx = headers.get(col_name)
    if col_idx is None:
        # 부분 일치 폴백 — 공백/줄바꿈/대소문자 변형에 대응
        norm = col_name.replace(" ", "").replace("\n", "").lower()
        for h, idx in headers.items():
            h_norm = h.replace(" ", "").replace("\n", "").lower()
            if norm in h_norm or h_norm in norm:
                col_idx = idx
                break
    if col_idx is None:
        return None
    return sheet.cell_value(row, col_idx)


def _get_str(sheet, row: int, headers: dict[str, int], col_name: str) -> str | None:
    """셀 값을 문자열로 반환. 빈 문자열은 None으로 정규화."""
    val = _get_cell(sheet, row, headers, col_name)
    if val is None:
        return None
    # xlrd는 숫자 셀을 float으로 반환하므로 정수형 표현으로 변환
    if isinstance(val, float) and val == int(val):
        s = str(int(val))
    else:
        s = str(val).strip()
    return s if s else None


def _get_num(sheet, row: int, headers: dict[str, int], col_name: str) -> float | None:
    """셀 값을 float으로 반환. 변환 불가 시 None."""
    val = _get_cell(sheet, row, headers, col_name)
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _get_int(sheet, row: int, headers: dict[str, int], col_name: str) -> int | None:
    """셀 값을 int로 반환. 변환 불가 시 None."""
    val = _get_num(sheet, row, headers, col_name)
    return int(val) if val is not None else None


def _parse_date(val, datemode: int):
    """Excel 날짜 값(시리얼 float 또는 문자열)을 `datetime.date`로 변환.

    - float: xlrd.xldate_as_tuple을 이용해 시리얼 번호 변환
    - str:  %Y-%m-%d / %Y%m%d / %Y.%m.%d 포맷 순서로 시도
    변환 실패 시 None 반환 (예외를 외부로 전파하지 않음).
    """
    if val is None or val == "":
        return None

    if isinstance(val, float):
        # YYYYMMDD 형식 숫자 (예: 20260529.0) → 문자열 변환 후 파싱
        int_val = int(val)
        if 19000101 <= int_val <= 29991231:
            s = str(int_val)
            try:
                return datetime.strptime(s, "%Y%m%d").date()
            except ValueError:
                pass
        # Excel serial date (작은 숫자)
        try:
            dt_tuple = xlrd.xldate_as_tuple(val, datemode)
            return datetime(*dt_tuple).date()
        except Exception:
            return None

    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _extract_core_count(spec: str) -> int:
    """규격 문자열에서 심수(Core Count)를 추출한다.

    Examples:
        "4C x 35SQ"   → 4
        "1C x 300SQ"  → 1
        "3C+1C x 6SQ" → 3   (첫 번째 숫자)
        ""            → 1   (파싱 불가 시 단심 기본값)
    """
    m = re.match(r"(\d+)\s*[Cc]", spec.strip())
    return int(m.group(1)) if m else 1
